"""Tests for reporting_engine pure helpers + a build smoke test.
Run: pytest -q test_reporting_engine.py"""
import numpy as np
import pandas as pd
from openpyxl import load_workbook

import reporting_engine as eng


# --- parse_bool ---------------------------------------------------------------
def test_parse_bool():
    assert eng.parse_bool("yes") and eng.parse_bool(1) and eng.parse_bool(True)
    assert not eng.parse_bool("no") and not eng.parse_bool(0)
    assert eng.parse_bool(None) is True and eng.parse_bool(float("nan")) is True


# --- colour / text normalisation (the Excel int-colour trap) ------------------
def test_resolve_style_colour_coercion():
    cfg = eng.resolve_sheet_style({
        "table_start_row": "5", "header_font_color": 0,        # 000000 read as int
        "desc_font_color": 595959,                              # all-digit -> int
        "header_bg": "1F4E78", "title_bg_color": 0, "desc_text": float("nan"),
    })
    assert cfg["header_font_color"] == "000000"
    assert cfg["desc_font_color"] == "595959"
    assert cfg["header_bg"] == "1F4E78"
    assert cfg["title_bg_color"] == "000000"      # kept, padded, truthy
    assert cfg["desc_text"] == ""                  # NaN -> skipped downstream
    assert cfg["table_start_row"] == 5

def test_resolve_style_missing_bg_dropped():
    cfg = eng.resolve_sheet_style({"table_start_row": "5"})
    assert "title_bg_color" not in cfg and "desc_bg_color" not in cfg


# --- audit rows ---------------------------------------------------------------
def test_build_audit_rows():
    rows = eng.build_audit_rows({"company": "ROYAL", "tolerance_pct": 0.005, "tolerance_abs": 1000})
    d = dict(rows)
    assert ("Execution Metadata", "") in rows          # section header
    assert d["Company"] == "ROYAL"
    assert d["Tolerance %"] == "0.5%"
    assert d["Tolerance Abs"] == "1,000.00"


# --- TOC + ordering -----------------------------------------------------------
def _details():
    return pd.DataFrame({
        "Sheet ID": ["DQR01", "DQR02", "DQR03"],
        "Sheet Name": ["01_A", "02_B", "03_C"],
        "Sheet Title": ["A", "B", "C"], "Color": ["#D9534F", "#5BC0DE", "#F0AD4E"],
        "Fade Color": ["#F2DEDE", "#D9EDF7", "#FCF8E3"],
        "Category": ["Action", "Info", "Warning"], "Action": ["Fix", "Read", "Review"],
        "Order": [1, 2, 3],
    })

def test_build_toc_rows_exists_flag():
    rows = eng.build_toc_rows(_details(), existing_ids={"DQR01", "DQR03"})
    by_id = {r["sheet_id"]: r for r in rows}
    assert by_id["DQR01"]["exists"] and not by_id["DQR02"]["exists"]
    assert by_id["DQR01"]["fade_color"] == "F2DEDE"

def test_build_ordered_sheets_skips_missing_and_orders():
    frames = {"DQR03": pd.DataFrame({"x": [1]}), "DQR01": pd.DataFrame({"x": [1]})}
    specs = eng.build_ordered_sheets(_details(), {"DQR01": {}, "DQR03": {}}, frames)
    assert [s["tab_name"] for s in specs] == ["01_A", "03_C"]   # Order respected, DQR02 skipped


# --- pct column detection -----------------------------------------------------
def test_column_plan_marks_pct():
    df = pd.DataFrame({"GROSS_PAID_DIFF": [1.0], "GROSS_PAID_PCT": [0.01], "Name": ["x"]})
    kinds = [k for _, k in eng._column_plan(df, pct_cols={"GROSS_PAID_PCT"})]
    assert kinds == ["float", "pct", "other"]


# --- full single-pass build smoke test ----------------------------------------
def test_build_workbook_smoke(tmp_path):
    details = _details()
    raw = {"DQR01": {"table_start_row": "5", "title_text": "A", "header_bg": "1F4E78",
                     "sheet_name": "01_A", "font_name": "Arial"},
           "DQR02": {"table_start_row": "5", "title_text": "Recon", "sheet_name": "02_B",
                     "font_name": "Arial"}}
    recon = pd.DataFrame({"LOB": ["m", "p"], "GROSS_PAID_DIFF": [5000.0, 1.0],
                          "GROSS_PAID_PCT": [0.001, 0.02]})
    frames = {"DQR01": pd.DataFrame({"Policy": ["P1"], "Amt": [10.0]}), "DQR02": recon}
    out = str(tmp_path / "wb.xlsx")
    eng.build_report_workbook(
        out,
        eng.build_ordered_sheets(details, raw, frames),
        eng.build_toc_rows(details, set(frames)),
        eng.build_audit_rows({"company": "ROYAL"}),
        recon_config={"tolerance_abs": 1000.0, "tolerance_pct": 0.005},
    )
    wb = load_workbook(out)
    assert wb.sheetnames[:2] == ["Audit Trail", "TOC"]
    assert "02_B" in wb.sheetnames
    cf = list(wb["02_B"].conditional_formatting)
    assert len(cf) >= 1            # recon CF present on the recon sheet
    assert len(list(wb["01_A"].conditional_formatting)) == 0  # none on non-recon sheet
