"""Tests for excel_export. Run: pytest -q test_excel_export.py"""

import os
import datetime

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import excel_export as ee


def _df():
    return pd.DataFrame({
        "Policy": ["P1", "P2", "P3"],
        "Loss Date": pd.to_datetime(["2024-01-01", "2024-02-15", "2024-03-30"]),
        "Gross Amount": [1000.5, 2500.0, 99999.99],
    })


def _base_cfg(path, **over):
    cfg = {
        "sheet_name": "Sheet1", "title_text": "My Report", "desc_text": "A description",
        "font_name": "Arial", "header_bg": "1F4E78", "header_font_color": "FFFFFF",
        "title_font_color": "1F4E78", "desc_font_color": "595959",
        "title_bg_color": "FFFFFF", "desc_bg_color": "FFFFFF", "table_border_color": "F2F2F2",
        "table_start_row": "5", "show_gridlines": "False", "float_number_format": "#,##0",
    }
    cfg.update(over)
    cfg = ee.apply_config_defaults(cfg)
    cfg["file_path"] = path
    return cfg


# --- parse_bool / flag default ------------------------------------------------
def test_parse_bool_truthy():
    for v in [True, 1, "TRUE", "true", "Yes", "y", "1", "on"]:
        assert ee.parse_bool(v) is True

def test_parse_bool_falsy():
    for v in [False, 0, "FALSE", "false", "No", "n", "0", "off"]:
        assert ee.parse_bool(v) is False

def test_parse_bool_missing_defaults_true():
    assert ee.parse_bool(None) is True
    assert ee.parse_bool(float("nan")) is True
    assert ee.parse_bool("") is True
    assert ee.parse_bool("weird") is True
    assert ee.parse_bool(None, default=False) is False

def test_apply_formatting_defaults_true_when_absent():
    cfg = ee.apply_config_defaults({"table_start_row": "5"})
    assert cfg["apply_formatting"] is True

def test_apply_formatting_reads_no():
    cfg = ee.apply_config_defaults({"table_start_row": "5", "apply_formatting": "no"})
    assert cfg["apply_formatting"] is False


# --- helpers ------------------------------------------------------------------
def test_build_output_path():
    p = ee.build_output_path("Output", "ACME", "Report.xlsx", "20260624_1200")
    assert p == "..\\..\\Output\\ACME Report.xlsx 20260624_1200.xlsx"

def test_handle_empty_dataset():
    out = ee.handle_empty_dataset(pd.DataFrame(columns=["a", "b"]))
    assert len(out) == 1 and out.iloc[0, 0] == "No data available."

def test_clean_date_columns():
    df = pd.DataFrame({"d": pd.to_datetime(["2024-01-01", "2024-02-02"])})
    out = ee.clean_date_columns(df)
    assert isinstance(out["d"].iloc[0], datetime.date)


# --- xlsxwriter new-file path -------------------------------------------------
def test_new_file_formatted(tmp_path):
    fp = str(tmp_path / "out.xlsx")
    cfg = _base_cfg(fp)
    ee.generate_formatted_excel(ee.clean_date_columns(_df()), cfg)
    wb = load_workbook(fp)
    ws = wb["Sheet1"]
    assert ws["A1"].value == "My Report"
    assert ws["A5"].value == "Policy"          # header on row 5
    assert ws["A6"].value == "P1"              # data right below
    assert ws["C6"].value == 1000.5
    assert ws["C6"].number_format == "#,##0"   # column-level float format applied
    assert ws.freeze_panes == "A6"

def test_new_file_unformatted_is_correct(tmp_path):
    fp = str(tmp_path / "plain.xlsx")
    cfg = _base_cfg(fp, apply_formatting="FALSE")
    ee.generate_formatted_excel(ee.clean_date_columns(_df()), cfg)
    wb = load_workbook(fp)
    ws = wb["Sheet1"]
    # No title styling, header written by pandas on the start row, values intact
    assert ws["A5"].value == "Policy"
    assert ws["A6"].value == "P1"
    assert ws["C6"].value == 1000.5
    assert ws["A1"].value is None              # no title in plain mode


# --- openpyxl append path -----------------------------------------------------
def test_append_adds_second_sheet(tmp_path):
    fp = str(tmp_path / "multi.xlsx")
    ee.generate_formatted_excel(ee.clean_date_columns(_df()), _base_cfg(fp, sheet_name="First"))
    ee.generate_formatted_excel(ee.clean_date_columns(_df()), _base_cfg(fp, sheet_name="Second"))
    wb = load_workbook(fp)
    assert wb.sheetnames == ["First", "Second"]
    assert wb["Second"]["A5"].value == "Policy"
    assert wb["Second"]["A1"].value == "My Report"

def test_append_unformatted(tmp_path):
    fp = str(tmp_path / "multi2.xlsx")
    ee.generate_formatted_excel(ee.clean_date_columns(_df()), _base_cfg(fp, sheet_name="First"))
    ee.generate_formatted_excel(
        ee.clean_date_columns(_df()), _base_cfg(fp, sheet_name="Second", apply_formatting="no")
    )
    wb = load_workbook(fp)
    ws = wb["Second"]
    assert ws["A5"].value == "Policy"          # header present (pandas)
    assert ws["A1"].value is None              # but no styling


# --- empty dataset end-to-end -------------------------------------------------
def test_empty_dataset_exports(tmp_path):
    fp = str(tmp_path / "empty.xlsx")
    df = ee.handle_empty_dataset(pd.DataFrame(columns=["a", "b", "c"]))
    ee.generate_formatted_excel(df, _base_cfg(fp))
    wb = load_workbook(fp)
    assert wb["Sheet1"]["A6"].value == "No data available."
